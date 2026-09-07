import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

NON_DOCUMENT_MARK = "비서류"
WEIGHT_RE = re.compile(r"^\s*([\d]+(?:[.,]\d+)?)\s*(?:까지|〃)?\s*$")

EMS_DESTINATIONS = (
    ("DE", "Германия", "독일", 10),
    ("FR", "Франция", "프랑스", 20),
    ("ES", "Испания", "스페인", 30),
    ("GB", "Великобритания", "영국", 40),
    ("US", "США", "미국", 50),
    ("RU", "Россия", "러시아", 60),
    ("EU", "Европа (остальные)", "3지역", 70),
    ("UA", "Украина", "3지역", 80),
    ("TR", "Турция", "3지역", 90),
    ("KZ", "Казахстан", "3지역", 100),
    ("UZ", "Узбекистан", "3지역", 110),
    ("KG", "Кыргызстан", "3지역", 120),
    ("KR", "Корея", None, 130),
)

HEADER_FIRST_CELLS = {"(kg)", "kg", "（kg）"}

COLUMN_TITLES_EN = {
    "호주": "Australia",
    "브라질": "Brazil",
    "캐나다": "Canada",
    "중국": "China",
    "프랑스": "France",
    "독일": "Germany",
    "홍콩": "Hong Kong",
    "인도네시아": "Indonesia",
    "일본": "Japan",
    "말레이시아": "Malaysia",
    "뉴질랜드": "New Zealand",
    "필리핀": "Philippines",
    "러시아": "Russia",
    "싱가포르": "Singapore",
    "스페인": "Spain",
    "대만": "Taiwan",
    "태국": "Thailand",
    "영국": "United Kingdom",
    "미국": "United States",
    "베트남": "Vietnam",
    "1지역": "Zone 1",
    "2지역": "Zone 2",
    "3지역": "Zone 3",
    "4지역": "Zone 4",
}


def column_title(code):
    return COLUMN_TITLES_EN.get(code, code)


def upsert_rate_column(code, sort=0):
    from .models import EmsRateColumn

    title = column_title(code)
    column, created = EmsRateColumn.objects.get_or_create(
        code=code,
        defaults={"title": title, "sort": sort},
    )
    if column.title != title:
        column.title = title
        column.save(update_fields=["title"])
    return column


def seed_ems_destinations():
    from .models import EmsDestination

    needed_codes = {column for *_, column, _ in EMS_DESTINATIONS if column}
    columns = {}
    for sort, code in enumerate(sorted(needed_codes)):
        columns[code] = upsert_rate_column(code, sort=sort)
    for code, name, column_code, sort in EMS_DESTINATIONS:
        defaults = {"name": name, "sort": sort, "is_active": True}
        if column_code:
            defaults["rate_column"] = columns[column_code]
        destination, created = EmsDestination.objects.get_or_create(code=code, defaults=defaults)
        if not created and destination.rate_column_id is None and column_code:
            destination.rate_column = columns[column_code]
            destination.save(update_fields=["rate_column"])


def _cell_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_weight_step_grams(value):
    text = _cell_str(value).replace(" ", "")
    match = WEIGHT_RE.match(text)
    if not match:
        return None
    kg = Decimal(match.group(1).replace(",", "."))
    grams = (kg * Decimal(1000)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return int(grams)


def parse_price_krw(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
    else:
        text = _cell_str(value).replace(",", "").replace(" ", "")
        if not text:
            return None
        try:
            number = Decimal(text)
        except InvalidOperation:
            return None
    won = (number * Decimal(1000)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    if won < 0:
        return None
    return int(won)


def _is_country_header(cells):
    first = _cell_str(cells[0]).replace(" ", "").lower()
    return first in HEADER_FIRST_CELLS


def parse_ems_xlsx(file_obj) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=file_obj, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
    finally:
        workbook.close()

    start = None
    for index, row in enumerate(rows):
        joined = " ".join(_cell_str(cell) for cell in row)
        if NON_DOCUMENT_MARK in joined:
            start = index
            break
    if start is None:
        raise ValueError("В файле нет блока 비서류 (недокументы).")

    columns = {}
    current_headers = None
    for row in rows[start + 1:]:
        cells = list(row) + [None] * max(0, 9 - len(row))
        if _is_country_header(cells):
            current_headers = [_cell_str(cell) for cell in cells[1:] if _cell_str(cell)]
            continue
        if not current_headers:
            continue
        weight = parse_weight_step_grams(cells[0])
        if weight is None:
            current_headers = None
            continue
        for offset, code in enumerate(current_headers):
            price = parse_price_krw(cells[offset + 1] if offset + 1 < len(cells) else None)
            if price is None:
                continue
            columns.setdefault(code, {})[weight] = price

    if not columns:
        raise ValueError("Не удалось прочитать ставки 비서류.")

    return {
        "columns": columns,
        "column_count": len(columns),
        "rate_count": sum(len(steps) for steps in columns.values()),
    }


def import_ems_xlsx(file_obj) -> dict:
    from django.db import transaction
    from django.utils import timezone
    from .models import EmsRate

    parsed = parse_ems_xlsx(file_obj)
    now = timezone.now()
    with transaction.atomic():
        for sort, code in enumerate(parsed["columns"]):
            column = upsert_rate_column(code, sort=sort)
            EmsRate.objects.filter(column=column).delete()
            EmsRate.objects.bulk_create(
                [
                    EmsRate(column=column, weight_grams=weight, price_krw=price)
                    for weight, price in sorted(parsed["columns"][code].items())
                ]
            )
            column.updated_at = now
            column.save(update_fields=["updated_at"])

        seed_ems_destinations()

    sample = None
    russia = parsed["columns"].get("러시아", {})
    if 1000 in russia:
        sample = {"column": column_title("러시아"), "weight_grams": 1000, "price_krw": russia[1000]}

    return {
        **parsed,
        "sample": sample,
    }


def ems_price_krw(destination_code, weight_grams):
    from .models import EmsDestination

    if weight_grams is None or weight_grams < 0:
        return None
    destination = (
        EmsDestination.objects.select_related("rate_column")
        .filter(code=destination_code)
        .first()
    )
    if destination is None or destination.rate_column_id is None:
        return None
    rate = (
        destination.rate_column.rates.filter(weight_grams__gte=weight_grams)
        .order_by("weight_grams")
        .first()
    )
    return None if rate is None else rate.price_krw
